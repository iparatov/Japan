"""Сборка офлайн-путеводителя: прототип Claude Design + таблица → dist/.

python3 build.py [src/table.xlsx]
"""
import base64, gzip, hashlib, json, re, sys, urllib.parse
from pathlib import Path

ROOT = Path(__file__).parent
PROTO = ROOT / 'src/prototype.html'
DIST = ROOT / 'dist'
APP = 'Trip%20App.dc.html'


# ---------- бандл ----------
def unbundle(html):
    man = json.loads(re.search(r'<script type="__bundler/manifest">(.*?)</script>', html, re.S).group(1))
    ext = json.loads(re.search(r'<script type="__bundler/ext_resources">(.*?)</script>', html, re.S).group(1))
    files = {}
    for uid, v in man.items():
        b = base64.b64decode(v['data'])
        if v.get('compressed') in (True, 'True'): b = gzip.decompress(b)
        files[uid] = b
    names = {e['id'].lstrip('./'): e['uuid'] for e in ext}
    return man, names, files


def rebundle(html, man, files):
    for uid, b in files.items():
        man[uid]['data'] = base64.b64encode(gzip.compress(b, 9, mtime=0)).decode()
        man[uid]['compressed'] = True
    m = re.search(r'(<script type="__bundler/manifest">)(.*?)(</script>)', html, re.S)
    return html[:m.start(2)] + json.dumps(man, separators=(',', ':')) + html[m.end(2):]


def patch(src, old, new, count=1):
    n = src.count(old)
    assert n == count, f'патч не применился ({n} вхождений): {old[:80]!r}'
    return src.replace(old, new)


# ---------- данные ----------
def hyperlink(v):
    """Текст и ссылка из ячейки (формула HYPERLINK или обычный текст)."""
    if v is None: return '', None
    v = str(v)
    m = re.match(r'=HYPERLINK\("(.*)"\s*[,;]\s*"((?:[^"]|"")*)"\)\s*$', v, re.S)
    return (m.group(2).replace('""', '"'), m.group(1).replace('"&"', '').replace('""', '"')) if m else (v.strip(), None)


def book_kind(t):
    t = (t or '').lower()
    if re.match(r'(нет|стол не бронир|на завтрак нет|не подтвержд|стол — нет)', t): return 'none'
    if 'только курс' in t: return 'course'
    if re.search(r'онлайн|tabelog|hot pepper|сайт|приложен', t): return 'online'
    if 'телефон' in t: return 'phone'
    return 'online' if t.startswith('да') else 'none'


def walk(t):
    m = re.match(r'\s*(\d+)\s*мин от (.+?)\*?\s*$', t or '')
    return (int(m.group(1)), m.group(2)) if m else (None, None)


def rest_name(n):
    n = re.sub(r'^\s*↳\s*', '', n)
    return re.sub(r'(\s+—\s+|,\s*)план Б к .*$', '', n).strip()


def same_place(a, b):
    a, b = a.lower(), b.lower()
    return a.split()[0] in b or b.split()[0] in a


def rest_group(rows):
    """Основной ресторан и запасные для одного приёма пищи.
    Запасной с «В базовый маршрут» = TRUE заменил прежний основной — тот убираем."""
    promoted = [r for r in rows if r.get('base') is True]
    main = promoted[0] if promoted else next((r for r in rows if not r['alt'] and r['timed']), rows[0])
    old = [r for r in rows if not r['alt'] and r['timed'] and r is not main]
    mname = rest_name(main['name'])
    backups = []
    for r in rows:
        if r is main or r in old: continue
        b = {'name': rest_name(r['name']), 'book': book_kind(r['book'])}
        n, ref = walk(r.get('walk'))
        if n and same_place(ref, mname): b['walk'] = n  # от прежнего основного — врёт, не показываем
        for k in ('rating', 'price', 'bookText', 'about', 'hours', 'how', 'mapUrl', 'tabelogUrl', 'bookStatus'):
            if r.get(k): b[k] = r[k]
        backups.append(b)
    m = {'name': mname}
    for k in ('rating', 'price', 'mapUrl', 'tabelogUrl'):
        if main.get(k): m[k] = main[k]
    if main.get('book'): m['book'] = main['book']
    return {'main': m, 'backups': backups}


def price(t):
    """Чек на человека: одна вилка, по отзывам — если есть."""
    t = t or ''
    rx = r'(до\s+)?¥\s?[\d\s]+(?:\s*[–-]\s*[\d\s]+)?(?:\s*тыс\.?)?'
    i = t.lower().find('отзыв')
    m = re.search(rx, t[i:]) if i >= 0 else None
    m = m or re.search(rx, t)
    if not m: return ''
    v = re.sub(r'\s+', ' ', m.group(0)).strip()
    if 'тыс' in v:
        v = re.sub(r'(\d+)', lambda n: f'{int(n.group(1)) * 1000:,}'.replace(',', ' '), v.replace('тыс.', '').replace('тыс', '')).strip()
    return v.replace(' – ', '–').replace('¥ ', '¥')


RISKY = re.compile(r'риск|гарант|бронь не сделана|бронь возможна|не обеспечен|пока не|очеред|нельзя|не помещ|не подтвержд|нерегуляр', re.I)
HOURS = re.compile(r'^(работает|открыт|в (субботу|воскресенье|праздник|понедельник|будни)|(понедельник|вторник|среда|четверг|пятница|суббота|воскресенье)\b|\d{1,2}:\d\d)', re.I)


def about_hours(t):
    """Описание и часы из «Почему риск / условия» — без самих рисков."""
    about, hours = [], []
    for sen in re.split(r'(?<=[.!])\s+', (t or '').strip()):
        sen = sen.strip()
        if not sen or RISKY.search(sen): continue
        (hours if HOURS.match(sen) and re.search(r'\d:\d\d', sen) else about).append(sen.rstrip('.'))
    return '. '.join(about) + ('.' if about else ''), '; '.join(re.sub(r'^Работает\s+', '', h) for h in hours)


def read_book(xlsx):
    import openpyxl
    wf = openpyxl.load_workbook(xlsx)
    wv = openpyxl.load_workbook(xlsx, data_only=True)
    return wf, wv


def cell(wf, wv, sheet, r, c):
    f = wf[sheet].cell(r, c)
    txt, url = hyperlink(f.value)
    if url is None:
        v = wv[sheet].cell(r, c).value
        txt = '' if v is None else str(v).strip()
        if f.hyperlink: url = f.hyperlink.target
    return txt, url


def read_restaurants(wf, wv):
    S = 'Рестораны — Tabelog'
    ws = wf[S]
    hdr = {str(ws.cell(5, c).value).strip(): c for c in range(1, ws.max_column + 1) if ws.cell(5, c).value}
    C = lambda r, name: cell(wf, wv, S, r, hdr[name])
    groups = {}
    for r in range(6, ws.max_row + 1):
        name, _ = C(r, 'Ресторан / кафе')
        ref, _ = C(r, 'Строка маршрута')
        if not name or not re.match(r'D\d+', ref): continue
        book, book_url = C(r, 'Возможность брони')
        base, _ = C(r, 'В базовый маршрут')
        about, hours = about_hours(C(r, 'Почему риск / важные условия')[0])
        rating = C(r, 'Рейтинг Tabelog')[0].replace(',', '.')
        row = dict(ref=int(ref[1:]), name=name, alt=name.lstrip().startswith('↳'),
                   base={'TRUE': True, 'FALSE': False}.get(base.upper()),
                   timed=bool(re.match(r'\d\d:\d\d', C(r, 'Время в плане')[0])),
                   book=book, bookText=book, walk=C(r, 'Пешком от основного')[0] if 'Пешком от основного' in hdr else '',
                   rating=rating if re.match(r'\d', rating) else '', price=price(C(r, 'Средний чек, ¥ / чел.')[0]),
                   about=about, hours=hours, how=C(r, 'Что сделать / влияние альтернативы')[0] if name.lstrip().startswith('↳') else '',
                   mapUrl=maps_url(C(r, 'Google Maps')[1], name), tabelogUrl=C(r, 'Tabelog')[1])
        if 'только курс' in book.lower() and '17:00' in book: row['bookStatus'] = 'Только курс после 17:00'
        groups.setdefault(row['ref'], []).append(row)
    return {ref: rest_group(rows) for ref, rows in groups.items()}


DASH = {'', '—', '-', 'None'}


class Copy:
    """Человеческие тексты: словарь «как в таблице → как в приложении».
    Пустая строка — убрать. Незнакомые тексты проходят как есть и попадают в отчёт."""
    def __init__(self, table, keep=()):
        self.table, self.keep, self.unknown = table, set(keep), []

    def __call__(self, t, track=True):
        t = (t or '').strip()
        if t in DASH: return ''
        if t in self.table: return self.table[t]
        if track and t not in self.keep and t not in self.unknown: self.unknown.append(t)
        return t


def clean_title(t):
    t = re.sub(r'^ПЛАН Б\s*—\s*', '', t.strip())
    t = re.sub(r'\s*—\s*план А$', '', t)
    t = re.sub(r'^(Ужин|Обед|Завтрак):\s*', r'\1 — ', t)
    return t


R = 'Базовый маршрут'
COL = dict(date=1, time=2, sub=3, title=4, line=6, meal=7, note=8, desc=9, link=10, tnotes=11, hours=14)
MEALS = ('Ужин', 'Обед', 'Завтрак', 'Кофе', 'Перекус', 'Покупка еды')


def route_rows(wf, wv):
    rows = {}
    for r in range(2, wf[R].max_row + 1):
        v = {k: cell(wf, wv, R, r, c) for k, c in COL.items()}
        if not v['date'][0]: continue
        rows[r] = v
    return rows


def classify(v):
    meal, title = v['meal'][0], v['title'][0]
    if any(m in meal for m in MEALS): return 'food'
    if 'Сон' == meal: return 'sleep'
    if meal.startswith('Подготовка'): return 'prep'
    return 'move' if '→' in title else 'place'


def norm(t):
    return re.sub(r'\W+', '', clean_title(t).lower())


def find_row(it, rows, date, used):
    """Строка таблицы для пункта: сначала по названию в тот же день, иначе та же строка, если тип не поменялся.
    Отметки привязаны к id пункта, поэтому вставка строк в таблицу их не сдвигает."""
    day = [r for r, v in rows.items() if v['date'][0][:5] == date and r not in used]
    same = [r for r in day if norm(rows[r]['title'][0]) == norm(it['title'])]
    r0 = int(it['id'][1:]) if re.match(r'r\d+$', it['id']) else None
    if r0 in same: return r0
    if same: return same[0]
    # ponytail: переименование + сдвиг строки одновременно не распознаём — такой пункт уйдёт в отчёт
    if r0 in day and classify(rows[r0]) == it['type']: return r0
    return None


def maps_url(url, name):
    """Ссылка для кнопок «На карте» / «Маршрут»: только Google Maps, в виде, который открывает iOS-приложение.
    maps/place/?q=place_id:… приложение ищет как текст — переводим в Maps URLs API. Не карта (сайт перевозчика,
    поиск билетов) — не кнопка «Маршрут», убираем."""
    if not url: return None
    m = re.search(r'[?&]q=place_id:([\w-]+)', url)
    if m: return 'https://www.google.com/maps/search/?api=1&query=%s&query_place_id=%s' % (urllib.parse.quote(name.split(' — ')[-1]), m.group(1))
    return url if re.match(r'https://(www\.google\.com/maps|maps\.app\.goo\.gl/)', url) else None


def item_from_row(r, v, it, copy, report):
    """Обновляет изменчивые поля пункта из строки таблицы; структуру (тип, план Б) берёт из прототипа."""
    it = dict(it) if it else {'id': 'r%d' % r, 'type': classify(v)}
    m = re.match(r'(\d\d:\d\d)(?:–(\d\d:\d\d))?', v['time'][0])
    if m: it['s'], it['e'] = m.group(1), m.group(2) or ''
    it['title'] = copy(clean_title(v['title'][0]), False) or clean_title(v['title'][0])
    for k in ('note', 'desc', 'tnotes', 'hours'):
        t = copy(v[k][0])
        if t: it[k] = t
        else: it.pop(k, None)
    url = maps_url(v['link'][1] or v['title'][1], clean_title(v['title'][0]))
    key = 'routeUrl' if it['type'] == 'move' else 'mapUrl'
    if url: it[key] = url
    else: it.pop(key, None)  # не оставлять ссылку из прототипа
    if it['type'] == 'move' and v['line'][0] not in DASH: it['line'] = v['line'][0]
    return it


def attach_restaurants(days, rows, groups, copy, report, row_of):
    by_id = {it['id']: it for d in days for it in d['items']}
    by_row = {r: by_id[i] for i, r in row_of.items() if i in by_id}  # строка таблицы → пункт (и для новых пунктов с id n…)
    for ref, g in groups.items():
        target = by_row.get(ref)
        if target is None and rows.get(ref, {}).get('title', ('',))[0].startswith('ПЛАН Б'):
            nxt = [r for r in rows if r > ref and rows[r]['date'][0] == rows[ref]['date'][0] and classify(rows[r]) == 'food']
            target = by_row.get(nxt[0]) if nxt else None
            if target is not None:  # запасной ужин из плана Б → в список запасных основного
                # время пешком у них считано от ресторана плана Б, а не от основного — не показываем
                g = {'main': None, 'backups': [{k: v for k, v in b.items() if k != 'walk'} for b in
                     [dict(g['main'], book=book_kind(g['main'].get('book')), bookText=g['main'].get('book', ''))] + g['backups']]}
        if target is None:
            report.append(f'Рестораны: строка маршрута D{ref} не найдена — пропущено'); continue
        if g['main']:
            m = g['main']
            target['mainName'] = m['name']
            target['rest'] = {k: copy(m[k]) if k == 'book' else m[k] for k in ('name', 'rating', 'price', 'book') if m.get(k)}
            if m.get('tabelogUrl'): target['rest']['tabelogUrl'] = m['tabelogUrl']
        seen = {b['name'] for b in target.get('backups', [])} if not g['main'] else set()
        bs = [] if g['main'] else list(target.get('backups', []))
        for b in g['backups']:
            if b['name'] in seen: continue
            seen.add(b['name'])
            b = {k: (copy(x) if k in ('bookText', 'about', 'how', 'hours') else x) for k, x in b.items()}
            bs.append({k: x for k, x in b.items() if x not in ('', None)})
        if bs: target['backups'] = bs
        else: target.pop('backups', None)


def build_days(days, wf, wv, copy, report):
    rows = route_rows(wf, wv)
    used = set()
    for d in days:
        new = []
        for it in d['items']:
            r = find_row(it, rows, d['date'], used)
            if r is None:
                report.append(f'Маршрут: пункт «{it["title"]}» ({it["id"]}) больше не найден в таблице — убран'); continue
            used.add(r)
            new.append((r, item_from_row(r, rows[r], it, copy, report)))
        d['_new'] = new
        d['items'] = [it for _, it in new]
        for a in (d.get('alt') or {}).get('items', []):
            r = int(a['id'][1:]); used.add(r)
            if r in rows: a['note'] = copy(rows[r]['note'][0]) or a.get('note', '')
        for it in d['items']:
            pb = (it.get('planB') or {}).get('alt')
            if pb:
                for r, v in rows.items():
                    if v['date'][0][:5] == d['date'] and clean_title(v['title'][0]) == pb.get('title'):
                        used.add(r); pb['note'] = copy(v['note'][0]); pb['hours'] = copy(v['hours'][0]) or pb.get('hours', '')
        rmap = {it['id']: r for r, it in new}  # id пункта → его строка сейчас (строки в таблице сдвигаются)
        for it in d['items'] + (d.get('alt') or {}).get('items', []):
            for o in (it, it.get('planB') or {}, (it.get('planB') or {}).get('alt') or {}, it.get('booking') or {}):
                for k in ('cond', 'text', 'desc'):
                    if isinstance(o.get(k), str): o[k] = copy(o[k], False)
            if (it.get('planB') or {}).get('alt', {}).get('title') and re.search(re.escape(it['planB']['alt']['title']) + r'[^.]*исключ', it.get('note', '') + ' '.join(v[0] for v in rows.get(rmap.get(it['id']), {}).values() if isinstance(v, tuple))):
                report.append(f'План Б «{it["planB"]["alt"]["title"]}» у «{it["title"]}» исключён в таблице — убран'); it.pop('planB')
    ids = {it['id'] for d in days for _, it in d['_new']}
    row_of = {}
    for d in days:
        for r, v in rows.items():
            if r in used or v['date'][0][:5] != d['date'] or v['title'][0].startswith('ПЛАН Б') or not re.match(r'\d\d:\d\d', v['time'][0]):
                continue
            it = item_from_row(r, v, None, copy, report)
            it['id'] = 'n' + hashlib.sha1((d['date'] + norm(v['title'][0])).encode()).hexdigest()[:8]
            if it['id'] in ids: continue
            ids.add(it['id']); d['_new'].append((r, it))
            report.append(f'Маршрут: новый пункт {d["date"]} «{it["title"]}» (тип: {it["type"]}) — проверьте, как он выглядит')
        row_of.update({it['id']: r for r, it in d['_new']})
        d['items'] = [it for _, it in sorted(d.pop('_new'), key=lambda x: x[0])]
        d['total'] = len(d['items'])
    attach_restaurants(days, rows, read_restaurants(wf, wv), copy, report, row_of)
    merge_cond(days)
    card_text(days, report)
    for d in days:  # ссылки-кнопки переходов: только на билет или расписание, не на поиск
        for it in d['items']:
            if 'trip.com/trains/asia/list' in (it.get('link') or {}).get('url', ''): it.pop('link')
            if it['title'] == 'Kintetsu Kyoto → Kintetsu-Nara':
                it['link'] = {'label': 'Открыть Kintetsu e-ticket', 'url': 'https://www.ticket.kintetsu.co.jp/vs/en/T/TZZ/TZZ10.do?op=tDisplayVisitorMenu#vs-ticketinfo'}
    return days


def card_text(days, report):
    """Чистка текстов карточек по правилу «Важно» (22.10): src/card_text.json — «дата | название» → поле → [было, стало].
    Правка применяется, только если текст в данных всё ещё «было»; если таблица поменялась — в отчёт, текст из таблицы."""
    edits = json.loads((ROOT / 'src/card_text.json').read_text(encoding='utf-8'))
    seen = set()
    for d in days:
        for it in d['items'] + (d.get('alt') or {}).get('items', []):
            key = d['date'] + ' | ' + it['title']
            for f, (was, now) in edits.get(key, {}).items():
                seen.add(key)
                if (it.get(f) or '') != was:
                    report.append(f'Тексты: «{key}», поле {f} — в таблице другой текст, правка не применена'); continue
                if now: it[f] = now
                else: it.pop(f, None)
    for key in edits.keys() - seen:
        report.append(f'Тексты: пункт «{key}» не найден — правка не применена')


def merge_cond(days):
    """«Учесть» (cond) больше не отдельная строка: вливаем в «Важно» (note), без повторов."""
    for d in days:
        for it in d['items'] + (d.get('alt') or {}).get('items', []):
            c = (it.pop('cond', '') or '').strip()
            if not c: continue
            n = (it.get('note') or '').strip()
            if norm(c) in norm(n): continue
            it['note'] = (n + ' ' + c).strip() if n else c


def check_wish(app, wf, wv, report):
    """Вишлист собран в прототипе вручную; сверяем с таблицей и сообщаем о расхождениях."""
    W = json.loads(re.search(r'^  static WISH = (.*);$', app, re.M).group(1))
    key = lambda x: re.sub(r'\W', '', str(x or '')).lower()
    have = {key(w.get('sku') or w['title']): w for w in W}
    S = 'Вишлист'
    for r in range(2, wf[S].max_row + 1):
        what, model, sku, yen = (cell(wf, wv, S, r, c)[0] for c in (2, 4, 9, 8))
        if not (what or model): continue
        w = have.get(key(sku or model or what))
        if w is None: report.append(f'Вишлист: в таблице новая позиция «{model or what}» — в приложении её нет'); continue
        y = re.sub(r'[^\d.]', '', yen)
        if y and float(y) != w.get('yen'): report.append(f'Вишлист: цена «{w["title"]}» в таблице ¥{yen}, в приложении ¥{w.get("yen")}')


# ---------- логика ----------
LOGIC = [
    # componentDidUpdate в рантайме вызывается без prevState: сохранение отметок падало
    ('  componentDidUpdate(pp, ps) {\n',
     '  componentDidUpdate(pp, ps) {\n    ps = ps || this._ps || {}; this._ps = this.state;\n'),
    # «Сейчас»: ночью — сон прошлого дня, в паузе — следующий пункт
    ('''  currentId(di) {
    if (di !== this.state.nowDay) return null;
    const list = this.items(di) || []; const now = this.toMin(this.state.nowTime);
    const hit = list.find(it => { if (!it.e) return false; const s = this.toMin(it.s), e = this.toMin(it.e); return e > s ? (now >= s && now < e) : now >= s; });
    return hit ? hit.id : null;
  }''',
     '''  nowRef() {
    const nd = this.state.nowDay, now = this.toMin(this.state.nowTime);
    if (nd < 0 || nd > 10) return null;
    const timed = di => (this.items(di) || []).filter(it => it.e && /^\\d/.test(it.s));
    const list = timed(nd);
    const hit = list.find(it => { const s = this.toMin(it.s), e = this.toMin(it.e); return e > s ? (now >= s && now < e) : now >= s; });
    if (hit) return { day: nd, id: hit.id, soon: false };
    if (!list.length || now < this.toMin(list[0].s)) {
      const prev = nd > 0 ? timed(nd - 1).find(it => this.toMin(it.e) < this.toMin(it.s) && now < this.toMin(it.e)) : null;
      if (prev) return { day: nd - 1, id: prev.id, soon: false };
    }
    const next = list.find(it => this.toMin(it.s) > now);
    return next ? { day: nd, id: next.id, soon: true } : null;
  }
  currentId(di) { const r = this.nowRef(); return r && r.day === di ? r.id : null; }'''),
    ('''  goNow() { const S = this.state, nd = S.nowDay; if (nd < 0 || nd > 10) return;''',
     '''  goNow() { const S = this.state, r = this.nowRef(); if (!r) return; const nd = r.day;'''),
    ("nowText: 'сейчас · ещё ' + minTxt(rem)",
     "nowText: (isCur && (this.nowRef() || {}).soon) ? 'скоро · через ' + minTxt(this.toMin(it.s) - now) : 'сейчас · ещё ' + minTxt(rem)"),
    ('      setTimeout(() => { const c = this.currentId(this.state.day); if (c) this.scrollToId(c, false, 220); }, 400);',
     '      { const r = this.nowRef(); if (r && r.day !== this.state.day) this.setState({ day: r.day }); }\n      setTimeout(() => { const c = this.currentId(this.state.day); if (c) this.scrollToId(c, false, 220); }, 400);'),
]


# ---------- тексты интерфейса ----------
UI_APP = [
    ('Отметки и кнопка «Сейчас» заработают в день прилёта. Ниже — первый день целиком.',
     'Отметки и кнопка «Сейчас» включатся в день прилёта. Ниже — первый день.'),
    ('title="Пункты этого дня в макет не перенесены"', 'title="На этот день планов нет"'),
    ('В макете заполнены 30.10, 31.10 и 04.11 — на них видны все типы карточек. В приложении здесь будет {{ dayTotal }} пунктов из таблицы.', 'Можно отдохнуть.'),
    ('Фильтр ничего не нашёл среди 13 броней поездки. Сбросьте его, чтобы увидеть весь список.', 'Таких броней нет.'),
    ('Свайп вбок или стрелки — соседние позиции · вниз — закрыть', 'Вбок — соседняя позиция, вниз — закрыть'),
    ('С иконки путеводитель открывается без Safari и работает без сети: маршрут, брони и вишлист хранятся на телефоне.',
     'С иконки путеводитель работает без интернета: маршрут, брони и вишлист сохранены на телефоне.'),
    ("attentionTitle: attn.length + ' брони не куплены',",
     "attentionTitle: attn.length === 1 ? '1 бронь не куплена' : attn.length + (attn.length < 5 ? ' брони не куплены' : ' броней не куплено'),"),
    (". Купите заранее — на сайте брони.'", ". Купить заранее на сайте брони.'"),
]
UI_CARD = [
    ("noteLabel: it.critical ? 'Важно:' : 'Примечание:',", "noteLabel: it.critical ? 'Важно:' : '',"),
]
CARD = 'Route%20Card.dc.html'

# Цвет — вариант C из Color Study (22.10): цвет города на выбранном дне, кольце прогресса,
# линии таймлайна и подсветке «сейчас»: Токио — --chart-1, Киото — --chart-5, Нара — --chart-3.
# Общее для вариантов: активные чипы — --primary, «Сделано» — --chart-2. Шапка дня — без подложки.
CITY = ['tokyo', 'tokyo', 'tokyo', 'tokyo', 'tokyo', 'tokyo', 'kyoto', 'kyoto', 'nara', 'tokyo', 'tokyo']  # по дням 30.10–09.11, как в Color Study
CITY_COLOR = {'tokyo': 'var(--chart-1)', 'kyoto': 'var(--chart-5)', 'nara': 'var(--chart-3)'}
# Плашка выбранного дня в светлой теме: подпись дня недели мелкая, нужен контраст ≥ 4,5:1.
# Белый на --chart-5 — 4,36:1, поэтому розовый чуть темнее (5,0:1); на оранжевом --chart-3 белый — 3,2:1,
# поэтому текст тёмный (4,8:1). В тёмной теме цвета светлые, текст --background проходит и так.
PILL_LIGHT = {'tokyo': 'var(--chart-1)', 'kyoto': 'color-mix(in oklab, var(--chart-5) 95%, black)', 'nara': 'var(--chart-3)'}
PILL_FG_LIGHT = {'tokyo': 'var(--background)', 'kyoto': 'var(--background)', 'nara': 'var(--foreground)'}
PALETTE_APP = [
    ('<div class="{{ themeClass }}" style="font-family:var(--font-sans);',
     '<div class="{{ themeClass }}" data-city="{{ cityKey }}" style="--city:{{ cityColor }}; --city-pill:{{ cityPill }}; --city-pill-fg:{{ cityPillFg }}; font-family:var(--font-sans);'),
    ("themeClass: dark ? 'dark' : '',",
     "themeClass: dark ? 'dark' : '', cityKey: %s[S.day] || 'tokyo', cityColor: %s[%s[S.day] || 'tokyo'],"
     " cityPill: dark ? %s[%s[S.day] || 'tokyo'] : %s[%s[S.day] || 'tokyo'],"
     " cityPillFg: dark ? 'var(--background)' : %s[%s[S.day] || 'tokyo'],"
     % (json.dumps(CITY), json.dumps(CITY_COLOR), json.dumps(CITY), json.dumps(CITY_COLOR), json.dumps(CITY),
        json.dumps(PILL_LIGHT), json.dumps(CITY), json.dumps(PILL_FG_LIGHT), json.dumps(CITY))),
    ('background:var(--foreground); transform:translateX', 'background:var(--city-pill); transform:translateX'),
    ("bg: selected ? 'var(--foreground)' : 'transparent',", "bg: selected ? 'var(--city-pill)' : 'transparent',"),
    ("fg: selected ? 'var(--background)' : 'var(--text-strong)',", "fg: selected ? 'var(--city-pill-fg)' : 'var(--text-strong)',"),
    ("wdFg: selected ? 'var(--background)'", "wdFg: selected ? 'var(--city-pill-fg)'"),
    ("dot: today ? (selected ? 'var(--background)'", "dot: today ? (selected ? 'var(--city-pill-fg)'"),
    ('style="stroke:var(--primary); stroke-width:6; stroke-linecap:{{ ringCap }}', 'style="stroke:var(--city); stroke-width:6; stroke-linecap:{{ ringCap }}'),
    ('style="stroke:var(--primary); stroke-width:6; stroke-linecap:{{ g.cap }}', 'style="stroke:var(--city); stroke-width:6; stroke-linecap:{{ g.cap }}'),
]
PALETTE_CARD = [  # --city приходит от корня приложения (CSS-переменная наследуется)
    ('width:1px; background:var(--line-hairline)', 'width:1px; background:var(--city, var(--line-hairline))'),
    ("const cardShadow = current ? 'inset 0 0 0 2px var(--primary), var(--shadow-popover)' : 'none';",
     "const cardShadow = current ? 'inset 0 0 0 2px var(--city, var(--primary)), var(--shadow-popover)' : 'none';"),
    ("nodeBg: current ? 'var(--primary)'", "nodeBg: current ? 'var(--city, var(--primary))'"),
    ("nodeRing: current ? 'var(--primary)'", "nodeRing: current ? 'var(--city, var(--primary))'"),
    ("moveRing: current ? 'var(--primary)'", "moveRing: current ? 'var(--city, var(--primary))'"),
    ("checkBg: done ? 'var(--foreground)' : 'transparent',", "checkBg: done ? 'var(--chart-2)' : 'transparent',"),
    ("checkRing: done ? 'var(--foreground)' : 'var(--line-field)',", "checkRing: done ? 'var(--chart-2)' : 'var(--line-field)',"),
]
PALETTE_DS = [  # чип фильтра (компонент дизайн-системы)
    ("background: active ? 'var(--foreground)' : 'var(--surface-card)',", "background: active ? 'var(--primary)' : 'var(--surface-card)',"),
    ("color: active ? 'var(--background)' : 'var(--text-body)',", "color: active ? 'var(--primary-foreground)' : 'var(--text-body)',"),
]


# ---------- карточки мест (22.10, правила «было / стало») ----------
# Порядок: описание → часы → «Важно» (!) → Сделано/Пропущено → На карте (+ⓘ) → остальные факты → План Б.
# Служебные строки «нужна сеть» / «Откроется в Google Maps» убраны; «Учесть» влито в «Важно» (merge_cond).
CARD_UI = [
 [
  "      netNote: offline ? 'Нет сети. Ссылки откроются, когда появится связь' : (it.link && !it.mapUrl && !it.routeUrl ? 'Откроется в браузере · нужна сеть' : 'Откроется в Google Maps · нужна сеть'),\n",
  ""
 ],
 [
  "<sc-if value=\"{{ hasDesc }}\"><p style=\"margin:0; font:var(--type-small); color:var(--text-body)\">{{ desc }}</p></sc-if>",
  "<sc-if value=\"{{ hasDesc }}\"><p style=\"margin:0; font:var(--type-small); color:var(--text-body)\">{{ desc }}</p></sc-if>\n          <sc-if value=\"{{ hasHours }}\"><div style=\"display:grid; grid-template-columns:16px auto minmax(0,1fr); column-gap:8px; align-items:baseline; font:var(--type-small)\"><x-import component-from-global-scope=\"DesignSystem_30197c.Icon\" name=\"clock\" size=\"{{ 14 }}\" style=\"{{ quietNudge }}\" hint-size=\"14px,14px\"></x-import><span style=\"color:var(--text-quiet)\">Часы работы</span><span style=\"color:var(--text-strong)\">{{ hours }}</span></div></sc-if>"
 ],
 [
  "              <span style=\"display:flex; align-items:center; gap:6px; font:var(--type-caption); color:var(--text-quiet)\">\n                <x-import component-from-global-scope=\"DesignSystem_30197c.Icon\" name=\"{{ netIcon }}\" size=\"{{ 13 }}\" hint-size=\"13px,13px\"></x-import>{{ netNote }}\n              </span>\n",
  "",
  2
 ],
 [
  "    if (it.hours) facts.push({ icon: 'clock', k: 'Часы работы', v: it.hours });\n",
  ""
 ],
 [
  "    if (it.cond) facts.push({ icon: 'info', k: 'Учесть', v: it.cond });\n",
  ""
 ],
 [
  "      noteLabel: it.critical ? 'Важно:' : '',\n      noteIcon: it.critical ? 'triangle-alert' : 'sticky-note',",
  "      noteLabel: 'Важно:',\n      noteIcon: it.critical ? 'triangle-alert' : 'circle-alert',\n      hours: it.hours, hasHours: !!it.hours, quietNudge: { color: 'var(--text-quiet)', transform: 'translateY(2px)' },"
 ]
]
APP_UI2 = [
 ["name=\"wifi\" size=\"{{ 13 }}\" hint-size=\"13px,13px\"></x-import>нужна сеть</span>", "name=\"wifi\" size=\"{{ 13 }}\" hint-size=\"13px,13px\"></x-import>онлайн</span>"],
 [
  ">На сайте ↗ · нужна сеть</x-import>",
  ">На сайте ↗</x-import>"
 ],
 [
  "            <span style=\"font:var(--type-caption); color:var(--text-quiet)\">Карта, WhatsApp и Tabelog открываются при наличии сети</span>\n",
  ""
 ],
 [
  "try { navigator.clipboard && navigator.clipboard.writeText(b.bookingNo); } catch (e) {}",
  "try { navigator.clipboard && navigator.clipboard.writeText(b.bookingNo).catch(() => {}); } catch (e) {}"
 ]
]

# ---------- брони ----------
# Данные — src/bookings.json (из писем и подтверждений); прототипный static BOOK заменяется целиком.
# Карточка: тип и название → когда → ячейки → номер брони (Скопировать) → Предъявить → Гости → Письмо → Важно → кнопки.
BOOK_UI = [
    ((ROOT / 'src/book_card_old.html').read_text(encoding='utf-8'), (ROOT / 'src/book_card.html').read_text(encoding='utf-8')),
    ("        cta: attention ? 'Открыть сайт брони' : b.kind === 'hotel' ? 'Открыть бронь' : (b.src === 'сайт' ? 'Открыть билет' : 'Открыть билет в ' + b.src),\n        net: offline ? 'Нет сети — откроется, когда появится связь' : (b.kind === 'hotel' || attention || b.src === 'сайт' ? 'Сайт брони' : b.src) + ' · нужна сеть' });",
     "        hasNo: !!b.bookingNo, noLabel: b.noLabel || 'Номер брони', copyAria: 'Скопировать номер брони ' + (b.bookingNo || ''), copyIcon: S.bCopied === b.id ? 'check' : 'copy',\n        senderAria: 'Скопировать отправителя: ' + (b.sender || '').split(' · ')[0], senderIcon: S.bCopied === b.id + ':s' ? 'check' : 'copy',\n        copySender: () => { try { navigator.clipboard && navigator.clipboard.writeText((b.sender || '').split(' · ')[0]).catch(() => {}); } catch (e) {} this.setState({ bCopied: b.id + ':s' }); clearTimeout(this.bct); this.bct = setTimeout(() => this.setState({ bCopied: null }), 1600); },\n        hasSeg: !!(b.segments && b.segments.length), segs: (b.segments || []).map(t => ({ t })), hasBag: !!b.baggage,\n        copy: () => { try { navigator.clipboard && navigator.clipboard.writeText(b.bookingNo); } catch (e) {} this.setState({ bCopied: b.id }); clearTimeout(this.bct); this.bct = setTimeout(() => this.setState({ bCopied: null }), 1600); },\n        hasShow: !!b.show, hasSeats: !!b.seats, hasSender: !!b.sender, hasImp: !!b.important, nudge: { transform: 'translateY(2px)', flex: 'none' },\n        hasUrl: !!b.url, hasMap: !!b.mapUrl, hasBtns: !!(b.url || b.mapUrl), btnCols: 'minmax(0,1fr)',  // кнопки друг под другом: «Открыть бронь в Booking» в полстроки не влезает\n        urlLabel: b.urlLabel || 'Открыть билет', openUrl: () => b.url && window.open(b.url, '_blank', 'noopener'), openMap: () => b.mapUrl && window.open(b.mapUrl, '_blank', 'noopener') });"),
    ("ticket: ['ticket', 'Билет'] };", "ticket: ['ticket', 'Билет'], flight: ['plane', 'Перелёт'] };"),
    ("const bChipDefs = [['all', 'Все'], ['train', 'Поезда'], ['bus', 'Автобусы'], ['hotel', 'Отели'], ['ticket', 'Билеты'], ['attention', 'Не куплено · ' + attn.length], ['confirm', 'Подтвердить']];",
     "const bChipDefs = [['all', 'Все'], ['flight', 'Перелёты'], ['train', 'Поезда'], ['bus', 'Автобусы'], ['hotel', 'Отели'], ['ticket', 'Билеты'], ['attention', 'Не куплено · ' + attn.length], ['confirm', 'Подтвердить']].filter(([k]) => k === 'all' || (k === 'attention' ? attn.length > 0 : k === 'confirm' ? B.some(b => b.st === 'confirm') : B.some(b => b.kind === k)));"),
    ("    const bGroups = [];\n    fl.forEach(b => {\n      let g = bGroups.find(x => x.day === b.day);\n      if (!g) { g = { day: b.day, label: Component.WD[b.day].toUpperCase() + ', ' + D[b.day].date, items: [], allPast: true }; bGroups.push(g); }",
     "    const bGroups = [], bPastGroups = [];\n    fl.forEach(b => {\n      const lst = isPast(b) ? bPastGroups : bGroups;\n      const label = b.groupLabel || Component.WD[b.day].toUpperCase() + ', ' + D[b.day].date;\n      let g = lst.find(x => x.label === label);\n      if (!g) { g = { day: b.day, label, items: [], allPast: true, isGroup: true }; lst.push(g); }"),
    ("cells: b.cells.map(([k, v]) => ({ k, v }))",
     "cells: b.cells.map(([k, v, d]) => ({ k: d ? k + ' · ' + d : k, v }))"),
    ("attention, ok: !attention,",
     "attention, ok: false,"),
    ("    bGroups.forEach(g => { g.opacity = g.allPast ? 'var(--opacity-disabled)' : 1; });",
     "    bGroups.forEach(g => { g.opacity = g.allPast ? 'var(--opacity-disabled)' : 1; });\n    // прошедшие брони — свёрнуты в конце списка\n    bPastGroups.forEach(g => { g.opacity = 'var(--opacity-disabled)'; });\n    const bPastN = bPastGroups.reduce((n, g) => n + g.items.length, 0);\n    if (bPastN) bGroups.push({ isToggle: true, label: (S.bPast ? 'Скрыть прошедшие · ' : 'Прошедшие · ') + bPastN, toggle: () => this.setState({ bPast: !S.bPast }) }, ...(S.bPast ? bPastGroups : []));"),
    ("bBoughtLine: 'Куплено ' + bBoughtN + ' из ' + B.length + ' броней',",
     "bBoughtLine: bBoughtN < B.length ? 'Куплено ' + bBoughtN + ' из ' + B.length : '',"),
    ("<span style=\"font:var(--type-caption); color:var(--text-quiet); padding:0 4px; font-variant-numeric:tabular-nums\">{{ bBoughtLine }}</span>",
     "<sc-if value=\"{{ bBoughtLine }}\"><span style=\"font:var(--type-caption); color:var(--text-quiet); padding:0 4px; font-variant-numeric:tabular-nums\">{{ bBoughtLine }}</span></sc-if>"),
    ("<sc-for list=\"{{ bGroups }}\" as=\"g\" hint-placeholder-count=\"3\">\n          <div style=\"display:flex; flex-direction:column; gap:10px; padding-top:8px\">",
     "<sc-for list=\"{{ bGroups }}\" as=\"g\" hint-placeholder-count=\"3\">\n          <sc-if value=\"{{ g.isToggle }}\"><button onClick=\"{{ g.toggle }}\" style=\"display:flex; align-items:center; justify-content:center; min-height:44px; margin-top:8px; border:none; border-radius:var(--r-control); background:var(--secondary); color:var(--text-body); font:600 15px/1.2 var(--font-sans); cursor:pointer\">{{ g.label }}</button></sc-if>\n          <sc-if value=\"{{ g.isGroup }}\"><div style=\"display:flex; flex-direction:column; gap:10px; padding-top:8px\">"),
    ("            </sc-for>\n          </div>\n        </sc-for>",
     "            </sc-for>\n          </div></sc-if>\n        </sc-for>"),
]


def bookings(days):
    """src/bookings.json → массив для static BOOK."""
    return json.loads((ROOT / 'src/bookings.json').read_text(encoding='utf-8'))['items']


def add_icons(files):
    """Иконки, которых нет в офлайн-наборе дизайн-системы (иначе компонент тянет их с unpkg — без сети пусто)."""
    import urllib.parse
    uid = next(u for u, b in files.items() if b.startswith(b'window.__DS_ICONS={'))
    s = files[uid].decode()
    extra = []
    for f in sorted((ROOT / 'src/icons').glob('*.svg')):
        if '"%s":' % f.stem in s[:200000]: continue
        svg = re.sub(r'<!--.*?-->', '', f.read_text(encoding='utf-8'), flags=re.S)
        svg = re.sub(r'\s+', ' ', svg).strip()
        extra.append('"%s":"data:image/svg+xml;utf8,%s"' % (f.stem, urllib.parse.quote(svg, safe='')))
    if extra: files[uid] = s.replace('window.__DS_ICONS={', 'window.__DS_ICONS={' + ','.join(extra) + ',', 1).encode()



# ---------- офлайн ----------
SW = """// Офлайн: всё приложение — один index.html, кладём его в кэш.
const CACHE = 'jp-trip-%s';
const FILES = ['./', './index.html', './manifest.webmanifest', './icon-192.png', './icon-512.png', './apple-touch-icon.png'];
self.addEventListener('install', e => e.waitUntil(caches.open(CACHE).then(c => c.addAll(FILES)).then(() => self.skipWaiting())));
self.addEventListener('activate', e => e.waitUntil(caches.keys().then(ks => Promise.all(ks.filter(k => k !== CACHE).map(k => caches.delete(k)))).then(() => self.clients.claim())));
self.addEventListener('fetch', e => {
  if (e.request.method !== 'GET' || new URL(e.request.url).origin !== location.origin) return;
  // Сначала кэш (работает без сети), в фоне — свежая версия к следующему запуску.
  e.respondWith(caches.open(CACHE).then(async c => {
    const hit = await c.match(e.request, { ignoreSearch: true }) || (e.request.mode === 'navigate' && await c.match('./index.html'));
    const net = fetch(e.request).then(r => { if (r.ok) c.put(e.request, r.clone()); return r; }).catch(() => hit || Response.error());
    return hit || net;
  }));
});
"""

MANIFEST = {
    'name': 'Япония · 30.10–09.11', 'short_name': 'Япония', 'lang': 'ru',
    'start_url': './', 'scope': './', 'display': 'standalone',
    'background_color': '#faf9f5', 'theme_color': '#1d6ad2',
    'icons': [{'src': 'icon-192.png', 'sizes': '192x192', 'type': 'image/png'},
              {'src': 'icon-512.png', 'sizes': '512x512', 'type': 'image/png', 'purpose': 'any maskable'}],
}

HEAD = """  <meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">
  <meta name="theme-color" content="#faf9f5" media="(prefers-color-scheme: light)">
  <meta name="theme-color" content="#141414" media="(prefers-color-scheme: dark)">
  <meta name="apple-mobile-web-app-capable" content="yes">
  <meta name="mobile-web-app-capable" content="yes">
  <meta name="apple-mobile-web-app-status-bar-style" content="default">
  <meta name="apple-mobile-web-app-title" content="Япония">
  <meta name="robots" content="noindex">
  <link rel="manifest" href="manifest.webmanifest">
  <link rel="apple-touch-icon" href="apple-touch-icon.png">
  <style>@media (prefers-color-scheme: dark) { body, #__bundler_thumbnail { background: #141414 !important; } }</style>
  <script>if ('serviceWorker' in navigator && (location.protocol === 'https:' || location.hostname === 'localhost')) navigator.serviceWorker.register('sw.js');</script>
"""


def icons():
    from PIL import Image, ImageDraw
    big = Image.new('RGB', (1024, 1024), '#1d6ad2')
    ImageDraw.Draw(big).ellipse((512 - 174, 512 - 174, 512 + 174, 512 + 174), fill='#ffffff')
    for name, n in (('icon-512.png', 512), ('icon-192.png', 192), ('apple-touch-icon.png', 180)):
        big.resize((n, n), Image.LANCZOS).save(DIST / name, optimize=True)


def build(xlsx='src/table.xlsx'):
    html = PROTO.read_text(encoding='utf-8')
    man, names, files = unbundle(html)
    app_uid = names[APP]
    app = files[app_uid].decode()
    for old, new in LOGIC + UI_APP + PALETTE_APP + BOOK_UI + APP_UI2:
        app = patch(app, old, new)
    ds_uid = next(u for u, b in files.items() if b.startswith(b'/* @ds-bundle'))
    ds = files[ds_uid].decode()
    for old, new in PALETTE_DS:
        ds = patch(ds, old, new)
    files[ds_uid] = ds.encode()
    files[app_uid] = app.encode()
    for n in (APP, CARD):  # ссылки на файлы дизайн-системы, которых нет в выгрузке: стили уже внутри бандла
        files[names[n]], k = re.subn(rb'<(link rel="stylesheet" href="(_ds/[^"]+|trip-theme\.css)">|script src="(_ds/[^"]+|icons-offline\.js)"></script>)\n?', b'', files[names[n]])
        assert k, f'в {n} не нашлось ссылок на дизайн-систему — проверьте выгрузку'
    app = files[app_uid].decode()
    card = files[names[CARD]].decode()
    for old, new, *n in UI_CARD + PALETTE_CARD + CARD_UI:
        card = patch(card, old, new, *n)
    files[names[CARD]] = card.encode()
    cj = json.loads((ROOT / 'src/copy.json').read_text(encoding='utf-8'))
    copy = Copy(cj.get('replace', {}), cj.get('keep', []))
    report = []
    m = re.search(r'^  static DAYS = (.*);$', app, re.M)
    wf, wv = read_book(ROOT / xlsx)
    days = build_days(json.loads(m.group(1)), wf, wv, copy, report)
    check_wish(app, wf, wv, report)
    app = app[:m.start(1)] + json.dumps(days, ensure_ascii=False, separators=(',', ':')) + app[m.end(1):]
    mb = re.search(r'^  static BOOK = \[\n.*?\n  \];$', app, re.M | re.S)
    assert mb, 'в прототипе не найден static BOOK'
    app = app[:mb.start()] + '  static BOOK = ' + json.dumps(bookings(days), ensure_ascii=False, separators=(',', ':')) + ';' + app[mb.end():]
    files[app_uid] = app.encode()
    add_icons(files)
    out = rebundle(html, man, files)
    out = patch(out, '  <title>', HEAD + '  <title>')
    DIST.mkdir(exist_ok=True)
    (DIST / 'index.html').write_text(out, encoding='utf-8')
    (DIST / 'sw.js').write_text(SW % hashlib.sha1(out.encode()).hexdigest()[:10], encoding='utf-8')
    (DIST / 'manifest.webmanifest').write_text(json.dumps(MANIFEST, ensure_ascii=False, indent=1), encoding='utf-8')
    icons()
    rep = ['# Отчёт сборки', ''] + ['- ' + r for r in report]
    if copy.unknown: rep += ['', '## Тексты без человеческой версии (добавьте в src/copy.json)', ''] + ['- ' + t for t in copy.unknown]
    (ROOT / 'report.md').write_text('\n'.join(rep) + '\n', encoding='utf-8')
    return out


if __name__ == '__main__':
    build(*sys.argv[1:2])
    print('dist/index.html готов')
